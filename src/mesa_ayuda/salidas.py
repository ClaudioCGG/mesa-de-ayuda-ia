"""
La bandeja de revisión: donde una persona lee cada borrador, lo corrige y lo
aprueba. Se genera en dos formatos:
    - Excel, para trabajar en el día a día (columna Estado con lista desplegable).
    - Página HTML, para mostrar el flujo (es la que se publica en GitHub Pages).
"""

from __future__ import annotations

import html
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ESTADOS = ["Pendiente de revisión", "Aprobado", "Editado y enviado", "Descartado"]


def _antecedente_txt(a) -> str:
    if not isinstance(a, dict):
        return ""
    return (f"{a['fecha']:%d/%m/%Y} · {a['persona']} · \"{a['asunto']}\" (hilo {a['hilo']})\n"
            f"Respuesta dada: {a['respuesta']}")


def tabla_bandeja(resultados: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame({
        "Fecha": resultados["fecha"].dt.strftime("%d/%m/%Y %H:%M"),
        "Persona": resultados["persona"],
        "Asunto": resultados["asunto"],
        "Consulta": resultados["consulta"],
        "Clasificación": resultados["clasificacion"],
        "Tono sugerido": resultados["tono"],
        "Similitud": resultados["similitud"].round(2),
        "Antecedente": resultados["antecedente"].map(_antecedente_txt),
        "Alertas": resultados["alertas"].map(lambda xs: "\n".join(xs)),
        "Borrador": resultados["borrador"],
        "Estado": resultados["estado"],
        "Comentario de quien revisa": "",
    })


def exportar_excel(resultados: pd.DataFrame, ruta: str | Path, evaluacion: pd.DataFrame | None = None) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    tabla = tabla_bandeja(resultados)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bandeja"
    ws.append(list(tabla.columns))
    for fila in tabla.itertuples(index=False):
        ws.append(list(fila))

    anchos = [16, 16, 28, 45, 18, 20, 10, 50, 40, 70, 20, 30]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1C5CAB")
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    colores = {"Repetida": "DCEAFB", "Posible antecedente": "FCEFC7", "Nueva": "D7F2E6"}
    for n, clasif in enumerate(tabla["Clasificación"], 2):
        ws.cell(n, 5).fill = PatternFill("solid", fgColor=colores.get(clasif, "FFFFFF"))

    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{len(tabla) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    if evaluacion is not None:
        we = wb.create_sheet("Evaluación")
        we.append(list(evaluacion.columns))
        for fila in evaluacion.itertuples(index=False):
            we.append([v if not isinstance(v, float) else round(v, 3) for v in fila])
        for celda in we[1]:
            celda.font = Font(bold=True)
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------
# Página HTML
# ----------------------------------------------------------------------
def _e(x) -> str:
    return html.escape(str(x))


CLASE_CHIP = {"Repetida": "rep", "Posible antecedente": "pos", "Nueva": "nue"}


def _tarjeta(r: pd.Series, i: int) -> str:
    a = r["antecedente"]
    chips = [f'<span class="chip {CLASE_CHIP[r["clasificacion"]]}">{_e(r["clasificacion"])}</span>',
             f'<span class="chip tono{" firme" if r["tono"] == "Firme y pedagógico" else ""}">Tono: {_e(r["tono"])}</span>']
    if r["clasificacion"] != "Nueva":
        chips.append(f'<span class="chip sim">Similitud {r["similitud"]:.2f}</span>')
    alertas = "".join(f"<li>{_e(x)}</li>" for x in r["alertas"])
    antecedente = ""
    if isinstance(a, dict):
        antecedente = (f'<div class="antecedente"><div class="etq">Antecedente citado</div>'
                       f'<div class="meta">{a["fecha"]:%d/%m/%Y} · {_e(a["persona"])} · “{_e(a["asunto"])}”</div>'
                       f'<p>{_e(a["respuesta"])}</p></div>')
    return f"""
<article class="tarjeta" data-clasif="{CLASE_CHIP[r['clasificacion']]}" data-firme="{int(r['tono'] == 'Firme y pedagógico')}" id="c{i}">
  <header>
    <div><strong>{_e(r['persona'])}</strong> <span class="fecha">{r['fecha']:%d/%m/%Y %H:%M}</span></div>
    <span class="estado" data-estado>Pendiente de revisión</span>
  </header>
  <h3>{_e(r['asunto'])}</h3>
  <blockquote>{_e(r['consulta'])}</blockquote>
  <div class="chips">{''.join(chips)}</div>
  {f'<ul class="alertas">{alertas}</ul>' if alertas else ''}
  {antecedente}
  <div class="etq">Borrador propuesto (editable)</div>
  <textarea rows="9" aria-label="Borrador para {_e(r['persona'])}">{_e(r['borrador'])}</textarea>
  <div class="acciones">
    <button class="btn primario" data-accion="Aprobado">Aprobar</button>
    <button class="btn" data-accion="Descartado">Descartar</button>
  </div>
</article>"""


def exportar_html(resultados: pd.DataFrame, ruta: str | Path, resumen: dict | None = None,
                  repo: str = "https://github.com/ClaudioCGG/mesa-de-ayuda-ia") -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    plantilla = (Path(__file__).with_name("plantilla_bandeja.html")).read_text(encoding="utf-8")
    tarjetas = "".join(_tarjeta(r, i) for i, r in resultados.iterrows())
    conteo = resultados["clasificacion"].value_counts().to_dict()
    s = resumen or {}
    reemplazos = {
        "{{TARJETAS}}": tarjetas,
        "{{REPO}}": repo,
        "{{N_TOTAL}}": str(len(resultados)),
        "{{N_REP}}": str(conteo.get("Repetida", 0)),
        "{{N_POS}}": str(conteo.get("Posible antecedente", 0)),
        "{{N_NUE}}": str(conteo.get("Nueva", 0)),
        "{{N_FIRME}}": str(int((resultados["tono"] == "Firme y pedagógico").sum())),
        "{{K_REP}}": f"{s.get('repetidas_con_antecedente_correcto', '–')} de {s.get('repetidas_reales', '–')}",
        "{{K_ERR}}": str(s.get("errores_con_confianza", "–")),
        "{{K_DES}}": f"{s.get('desorden_detectados', '–')} de {s.get('desorden_reales', '–')}",
        "{{K_REV}}": str(s.get("a_revision", "–")),
    }
    for k, v in reemplazos.items():
        plantilla = plantilla.replace(k, v)
    ruta.write_text(plantilla, encoding="utf-8")
    return ruta
