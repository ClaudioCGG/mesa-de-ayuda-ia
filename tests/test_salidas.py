from openpyxl import load_workbook

from mesa_ayuda.historial import cargar_mensajes, exportar_expedientes
from mesa_ayuda.salidas import exportar_excel, exportar_html

from conftest import RAIZ


def test_excel_tiene_bandeja_y_evaluacion(corrida, tmp_path):
    ruta = exportar_excel(corrida["resultados"], tmp_path / "b.xlsx", corrida["evaluacion"])
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Bandeja", "Evaluación"]
    ws = wb["Bandeja"]
    assert ws.max_row == len(corrida["resultados"]) + 1
    assert ws.data_validations.dataValidation  # lista desplegable de Estado


def test_html_escapa_texto_y_no_deja_marcadores(corrida, tmp_path):
    res = corrida["resultados"].copy()
    res.loc[0, "consulta"] = "<script>alert(1)</script>"
    html = exportar_html(res, tmp_path / "b.html", corrida["resumen"]).read_text(encoding="utf-8")
    assert "<script>alert(1)" not in html and "&lt;script&gt;" in html
    assert "{{" not in html
    assert html.count('class="tarjeta"') == len(res)


def test_expedientes_mismo_formato_que_apps_script():
    txt = exportar_expedientes(cargar_mensajes(RAIZ / "data" / "historial.csv"))
    assert txt.startswith("=== EXPEDIENTES DE MAILS")
    assert txt.count("HILO #") == 64
